"""Excel 内容検索「とりこぼし」調査テスト（調査フェーズ専用）。

目的: `SearchMode.EXCEL` でセルに存在するはずのキーワードがヒットしない
事象の原因を、再現テストで特定する。**本テストは調査用**であり、現状の
（=バグを含む）挙動を assert で固定して原因を可視化している。後の実装
フェーズで「あるべき挙動」への反転テストとして流用する前提。

対応レポート: `docs/investigation/excel_search_report.md`
フィクスチャ生成: `tests/fixtures/gen_excel_fixtures.py`

本体コード(app/ 以下)は一切変更していない。実験（reset_dimensions 等）は
テストコード側で完結させ、当たり判定はレポート用に記録する。
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent / "fixtures"))
import gen_excel_fixtures as gen  # noqa: E402

from app.engine.excel_reader import (  # noqa: E402
    may_contain_keyword, search_in_excel,
)
from app.engine.search_engine import (  # noqa: E402
    SearchMode, SearchOptions, SearchStats, search,
)


@pytest.fixture(scope="module")
def fx(tmp_path_factory):
    """全フィクスチャを一時ディレクトリに生成（gitignore 対象を汚さない）。"""
    d = tmp_path_factory.mktemp("excel_fixtures")
    return gen.generate_all(d)


def _hits(path, keyword):
    return list(search_in_excel(path, keyword))


# ---------------------------------------------------------------------------
# H1: プレフィルタが XML エンティティを未デコードのまま照合し、
#     &/</> を含むキーワードを誤って除外する
# ---------------------------------------------------------------------------
class TestH1EntityPrefilter:
    def test_ampersand_keyword_is_dropped(self, fx):
        p = fx["entity"]
        # セルには "R&D部門" が実在するが…
        # プレフィルタが False を返し本走査に到達しない → 取りこぼし
        assert may_contain_keyword(p, "R&D") is False
        assert _hits(p, "R&D") == []

    def test_angle_bracket_keyword_is_dropped(self, fx):
        p = fx["entity"]
        # "<重要>" も &lt; &gt; のまま照合され除外される
        assert may_contain_keyword(p, "<重要>") is False
        assert _hits(p, "<重要>") == []

    def test_quote_keyword_is_found(self, fx):
        p = fx["entity"]
        # 引用符は XML テキスト内で非エスケープ → これはヒットする（対照群）
        assert may_contain_keyword(p, '"引用"') is True
        assert _hits(p, '"引用"') == [("Sheet", "C1", '"引用"')]

    def test_plain_substring_of_entity_cell_is_found(self, fx):
        p = fx["entity"]
        # エンティティ文字を跨がない部分文字列なら拾える
        assert _hits(p, "部門") == [("Sheet", "A1", "R&D部門")]
        assert _hits(p, "重要") == [("Sheet", "B1", "<重要>")]


# ---------------------------------------------------------------------------
# H2: 日付/パーセント/桁区切り等の「表示書式文字列」は sharedStrings にも
#     セル値(str)にも現れず、表示どおりのキーワードでは一致しない
# ---------------------------------------------------------------------------
class TestH2DisplayFormat:
    # 表示どおりのキーワードは全て取りこぼす
    @pytest.mark.parametrize("kw", ["2026/07/04", "1,000", "15%"])
    def test_display_format_misses(self, fx, kw):
        assert _hits(fx["formats"], kw) == []

    # 内部表現寄りのキーワードなら一致する（str(datetime)/生数値）
    def test_internal_repr_hits(self, fx):
        p = fx["formats"]
        assert _hits(p, "2026-07-04") == [
            ("Sheet", "A1", "2026-07-04 00:00:00")]
        assert _hits(p, "1000") == [("Sheet", "B1", "1000")]
        assert _hits(p, "0.15") == [("Sheet", "C1", "0.15")]

    def test_comma_form_passes_prefilter_but_scan_misses(self, fx):
        # "1,000" は _NUMERIC_RE に一致しプレフィルタは通過するが、
        # 本走査で str(1000)=="1000" と照合され外れる（多段の取りこぼし）
        p = fx["formats"]
        assert may_contain_keyword(p, "1,000") is True
        assert _hits(p, "1,000") == []


# ---------------------------------------------------------------------------
# H3: read_only が dimension 宣言を信用し範囲外セルを取りこぼす
#     → openpyxl 3.1.5 でも再現する。dimension を実データより狭く偽ると、
#       read_only の iter_rows() は宣言範囲に丸められ範囲外セルを落とす。
# ---------------------------------------------------------------------------
class TestH3BadDimension:
    def test_out_of_dimension_cell_is_missed(self, fx):
        # dimension を A1:A1 に狭めると、E10 の実セルが取りこぼされる。
        # プレフィルタは通過している（worksheet XML には値が残る）ため、
        # 「XMLには在るのに本走査で見えない」典型的な取りこぼし。
        p = fx["bad_dimension"]
        assert may_contain_keyword(p, "TARGET_OUT_OF_DIM") is True
        assert _hits(p, "TARGET_OUT_OF_DIM") == []

    def test_reset_dimensions_recovers_the_cell(self, fx):
        # 実験: load 後に reset_dimensions() を挟むと範囲が再計算され、
        # E10 を読めるようになる（= 修正候補の裏付け）。
        p = fx["bad_dimension"]
        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        found = []
        for ws in wb.worksheets:
            ws.reset_dimensions()
            for r_i, row in enumerate(ws.iter_rows(values_only=True), 1):
                for c_i, val in enumerate(row, 1):
                    if val is not None and "TARGET_OUT_OF_DIM" in str(val):
                        found.append((r_i, c_i, val))
        wb.close()
        assert found == [(10, 5, "TARGET_OUT_OF_DIM")]


# ---------------------------------------------------------------------------
# H4: data_only=True でキャッシュ値のない数式セルは None → 取りこぼし
# ---------------------------------------------------------------------------
class TestH4FormulaNoCache:
    def test_formula_result_missed(self, fx):
        # =CONCATENATE("HIT","KEY") の結果 "HITKEY" は
        # キャッシュ無しのため data_only=True では取得できない
        p = fx["formula_nocache"]
        assert _hits(p, "HITKEY") == []
        # プレフィルタも通らない（XMLには "HIT" と "KEY" が分断されている）
        assert may_contain_keyword(p, "HITKEY") is False

    def test_formula_text_visible_only_without_data_only(self, fx):
        # data_only=False なら数式文字列そのものが読める（本体は data_only=True）
        p = fx["formula_nocache"]
        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=False)
        vals = [c.value for r in wb.worksheets[0].iter_rows()
                for c in r if c.value is not None]
        wb.close()
        assert vals == ['=CONCATENATE("HIT","KEY")']
        # 数式文字列 "CONCATENATE" はプレフィルタを通るが…
        assert may_contain_keyword(p, "CONCATENATE") is True
        # 本走査は data_only=True なので数式セルが None になり取りこぼす
        assert _hits(p, "CONCATENATE") == []


# ---------------------------------------------------------------------------
# H5: 開けないブック(暗号化/破損/オンライン専用)は silent skip。
#     stats.skipped にも計上されず診断不能。
# ---------------------------------------------------------------------------
class TestH5SilentSkip:
    def test_unopenable_file_yields_nothing_silently(self, fx):
        p = fx["fake_encrypted"]
        # プレフィルタは BadZipFile を握りつぶし True（本走査に委ねる）
        assert may_contain_keyword(p, "x") is True
        # load_workbook 失敗 → 空。例外も出ない
        assert _hits(p, "x") == []

    def test_engine_does_not_count_skip(self, fx, tmp_path):
        # 開けない .xlsx を含むツリーを search() で回しても skipped は増えない
        d = tmp_path / "tree"
        d.mkdir()
        (d / "broken.xlsx").write_bytes(b"not a zip")
        opts = SearchOptions(keyword="anything", modes={SearchMode.EXCEL})
        stats = SearchStats()
        hits = list(search(d, opts, stats=stats))
        assert hits == []
        assert stats.scanned == 1
        assert stats.skipped == 0  # ← 開けなかった事実がどこにも残らない


# ---------------------------------------------------------------------------
# H6: サイズ上限超過はスキップされ skipped に計上される（仕様）。
#     ただし理由(サイズ/読込失敗/プレフィルタ除外)の区別は UI に伝わらない。
# ---------------------------------------------------------------------------
class TestH6SizeLimit:
    def test_oversize_excel_counted_as_skipped(self, fx, tmp_path):
        d = tmp_path / "tree"
        d.mkdir()
        # 実ブックをコピー配置し、極小上限で必ず超過させる
        import shutil
        shutil.copy(fx["inline_string"], d / "big.xlsx")
        opts = SearchOptions(keyword="INLINE_UNIQUE_TOKEN",
                             modes={SearchMode.EXCEL}, max_file_size=10)
        stats = SearchStats()
        hits = list(search(d, opts, stats=stats))
        assert hits == []
        assert stats.skipped == 1  # サイズ理由。H5 と同じ skipped で区別不能

    def test_within_limit_excel_is_found(self, fx, tmp_path):
        d = tmp_path / "tree"
        d.mkdir()
        import shutil
        shutil.copy(fx["inline_string"], d / "ok.xlsx")
        opts = SearchOptions(keyword="INLINE_UNIQUE_TOKEN",
                             modes={SearchMode.EXCEL})  # 既定 50MB
        stats = SearchStats()
        hits = list(search(d, opts, stats=stats))
        assert len(hits) == 1
        assert hits[0].detail == "Sheet!A1: INLINE_UNIQUE_TOKEN"
        assert stats.skipped == 0


# ---------------------------------------------------------------------------
# Step3: プレフィルタ may_contain_keyword の網羅チェック（真値 vs 返り値）
# ---------------------------------------------------------------------------
class TestPrefilterMatrix:
    """実セル内容に対し may_contain_keyword が正しく True/False を返すか。

    False が返るのに実在する = 取りこぼしの原因。ここでは entity/formats/
    inline のセル内容を横断して代表的キーワード群を検証する。
    """

    def test_ascii_and_japanese_present(self, fx):
        p = fx["inline_string"]
        assert may_contain_keyword(p, "INLINE_UNIQUE_TOKEN") is True
        assert may_contain_keyword(p, "inline_unique_token") is True  # 小文字
        assert may_contain_keyword(p, "存在しない語") is False

    def test_entity_chars_falsely_filtered(self, fx):
        # & < > を含むと実在しても False（取りこぼし源）
        p = fx["entity"]
        assert may_contain_keyword(p, "R&D") is False   # 実在: R&D部門
        assert may_contain_keyword(p, "<重要>") is False  # 実在: <重要>
        # 引用符は素通り（非エスケープ）
        assert may_contain_keyword(p, '"引用"') is True

    def test_numeric_keywords_always_pass(self, fx):
        # _NUMERIC_RE に一致 → 数値セルの可能性ありとして常に True
        p = fx["formats"]
        for kw in ["123", "1.5", "-42", "1e5", "1,000"]:
            assert may_contain_keyword(p, kw) is True

    def test_numeric_with_symbol_falls_to_string_scan(self, fx):
        # 数値+記号は _NUMERIC_RE から外れ文字列走査へ。実在しなければ False
        p = fx["formats"]
        # "100円" "¥1,000" は formats に無い → 文字列走査で False
        assert may_contain_keyword(p, "100円") is False
        assert may_contain_keyword(p, "¥1,000") is False
        # "2026/07/04" も同様（内部はシリアル値）→ False = 日付取りこぼしの源
        assert may_contain_keyword(p, "2026/07/04") is False
